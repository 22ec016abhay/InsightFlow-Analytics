"""
report_generator.py
====================
Builds a professional, multi-sheet Excel report from analyzer.py's
results, using OpenPyXL. Each analysis (monthly sales, top customers,
best/worst products, store performance, stock alerts...) gets its own
formatted sheet: bold colored headers, auto-sized columns, filters,
and a totals row where it makes sense.
"""
 
import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))
 
from datetime import datetime
 
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
 
import config
from utils import get_logger, timer_decorator, log_errors
 
logger = get_logger(__name__)
 
# ------------------------------------------------------------------
# Shared styling constants
# ------------------------------------------------------------------
HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
TOTAL_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
TOTAL_FONT = Font(bold=True, size=11, name="Calibri")
THIN_BORDER = Border(*(Side(style="thin", color="CCCCCC"),) * 4)
 
 
def _write_dataframe(ws: Worksheet, df: pd.DataFrame, sheet_title: str,
                      numeric_totals: list[str] | None = None) -> None:
    """
    Write a DataFrame into a worksheet with professional formatting:
    bold header row, auto-sized columns, autofilter, and an optional
    totals row summing the given numeric columns.
    """
    # Title row
    ws["A1"] = sheet_title
    ws["A1"].font = Font(bold=True, size=14, name="Calibri", color="1E3A5F")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(df.columns), 1))
 
    header_row = 3
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name.replace("_", " ").title())
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
 
    # Data rows
    for row_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            if isinstance(value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
 
    last_data_row = header_row + len(df)
 
    # Totals row (sum of specified numeric columns)
    if numeric_totals:
        total_row = last_data_row + 1
        ws.cell(row=total_row, column=1, value="TOTAL").font = TOTAL_FONT
        for col_name in numeric_totals:
            if col_name in df.columns:
                col_idx = list(df.columns).index(col_name) + 1
                col_letter = get_column_letter(col_idx)
                cell = ws.cell(
                    row=total_row, column=col_idx,
                    value=f"=SUM({col_letter}{header_row + 1}:{col_letter}{last_data_row})",
                )
                cell.font = TOTAL_FONT
                cell.fill = TOTAL_FILL
                cell.number_format = "#,##0.00"
        for col_idx in range(1, len(df.columns) + 1):
            ws.cell(row=total_row, column=col_idx).fill = TOTAL_FILL
 
    # Autofilter on the header row
    last_col_letter = get_column_letter(len(df.columns))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"
 
    # Auto-size columns based on content width
    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max(
            [len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str).tolist()]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)
 
    ws.freeze_panes = f"A{header_row + 1}"
 
 
@timer_decorator
@log_errors
def generate_excel_report(analysis_results: dict, output_filename: str = "InsightFlow_Business_Report.xlsx") -> Path:
    """
    Build the full multi-sheet Excel workbook from analyzer.py's results
    and save it to reports/.
    """
    config.ensure_directories_exist()
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet
 
    # --- Summary sheet ---
    ws_summary = wb.create_sheet("Summary")
    ws_summary["A1"] = "InsightFlow Analytics — Business Report"
    ws_summary["A1"].font = Font(bold=True, size=16, color="1E3A5F")
    ws_summary["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws_summary["A2"].font = Font(italic=True, size=10, color="666666")
 
    aov = analysis_results["aov_stats"]
    inv = analysis_results["inventory_summary"]
    summary_lines = [
        ("Average Order Value", f"${aov['mean']:,.2f}"),
        ("Median Order Value", f"${aov['median']:,.2f}"),
        ("Total Units In Stock", f"{inv['total_units_in_stock']:,}"),
        ("Items Out Of Stock", inv["items_out_of_stock"]),
        ("Low-Stock Alerts", len(analysis_results["stock_alerts"])),
    ]
    row = 4
    for label, value in summary_lines:
        ws_summary.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=row, column=2, value=value)
        row += 1
    ws_summary.column_dimensions["A"].width = 26
    ws_summary.column_dimensions["B"].width = 20
 
    # --- One sheet per analysis ---
    sheet_plan = [
        ("Monthly Sales", analysis_results["monthly_sales"], ["total_revenue", "total_orders"]),
        ("Top Customers", analysis_results["top_customers"], ["total_spent"]),
        ("Best Products", analysis_results["best_products"], ["units_sold", "revenue"]),
        ("Worst Products", analysis_results["worst_products"], ["units_sold", "revenue"]),
        ("Category Performance", analysis_results["category_performance"], ["revenue", "profit"]),
        ("Store Performance", analysis_results["store_performance"], ["revenue", "profit"]),
        ("Regional Performance", analysis_results["regional_performance"], ["revenue", "profit"]),
        ("Stock Alerts", analysis_results["stock_alerts"], None),
    ]
 
    for sheet_name, df, totals in sheet_plan:
        ws = wb.create_sheet(sheet_name[:31])  # Excel sheet name limit
        _write_dataframe(ws, df.reset_index(drop=True), sheet_name, numeric_totals=totals)
 
    output_path = config.REPORTS_DIR / output_filename
    wb.save(output_path)
    logger.info(f"Excel report saved -> {output_path}")
    return output_path
 
 
if __name__ == "__main__":
    # Quick manual test: run "python src/report_generator.py"
    from analyzer import run_full_analysis
 
    results = run_full_analysis()
    path = generate_excel_report(results)
    print(f"\nExcel report generated: {path}")